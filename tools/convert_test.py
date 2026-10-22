"""PDF -> Word 真实转换自测（开发机用，不依赖 Docker）。

生成一份含标题、段落、表格的测试 PDF，调用与生产完全相同的
convert_worker.py，再用 python-docx 打开结果，检查可读性。

关键点：测试 PDF 必须**嵌入真正的中文字体**并带 ToUnicode 映射，
否则导出的是缺字占位符，测出来的“乱码”是夹具问题而不是转换器问题。

用法（包根目录，需已安装 requirements.txt + reportlab）：
    python tools/convert_test.py
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
WORKER = ROOT / "services" / "portal" / "app" / "convert_worker.py"

# 候选中文字体：(路径, TTC 子字体序号)
CJK_CANDIDATES = [
    (r"C:\Windows\Fonts\msyh.ttc", 0),
    (r"C:\Windows\Fonts\simhei.ttf", 0),
    (r"C:\Windows\Fonts\simsun.ttc", 0),
    ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 0),
    ("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc", 0),
    ("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc", 0),
    ("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 0),
    ("/System/Library/Fonts/PingFang.ttc", 0),
]

SAMPLE_TEXT = [
    ("title", "内部文件预览平台 转换测试报告"),
    ("body", "本文件用于验证 PDF 转 Word 链路的可读性，包括文字完整性、段落顺序与可编辑性。验收标准为内部可读，不要求版式完全还原。"),
    ("h2", "一、测试范围"),
    ("body", "纯文本段落识别；简单表格结构；中文字符编码。"),
    ("h2", "二、测试数据"),
]
TABLE_DATA = [
    ["序号", "项目", "结论"],
    ["1", "文字完整性", "通过"],
    ["2", "段落顺序", "通过"],
    ["3", "表格结构", "观察"],
]
CONCLUSION = "若转换结果中以上文字可选中、可编辑，且顺序正确，则满足“内部可读”的验收口径。"


def register_cjk() -> str | None:
    """注册一个可用中文字体，返回注册名；找不到则返回 None。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path, idx in CJK_CANDIDATES:
        if not Path(path).exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("CJK", path, subfontIndex=idx))
            return "CJK"
        except Exception:  # noqa: BLE001
            continue
    return None


def make_pdf(path: Path, font: str | None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base_font = font or "Helvetica"
    cjk = font is not None

    if not cjk:
        print("  [WARN] 未找到中文字体，改用纯 ASCII 夹具（中文断言将跳过）")

    def st(name, size, leading, bold=False, space=0):
        return ParagraphStyle(
            name,
            fontName=base_font,
            fontSize=size,
            leading=leading,
            spaceAfter=space,
        )

    title = st("T", 16, 22, space=10)
    body = st("B", 10.5, 18, space=6)
    h2 = st("H", 12, 18, space=6)

    text = SAMPLE_TEXT if cjk else [
        ("title", "PDF to Word Conversion Test Report"),
        ("body", "This fixture verifies readability after conversion: text completeness, paragraph order and editability."),
        ("h2", "1. Scope"),
        ("body", "Plain text extraction; simple table structure; character encoding."),
        ("h2", "2. Data"),
    ]
    rows = TABLE_DATA if cjk else [
        ["No", "Item", "Result"],
        ["1", "Text", "Pass"],
        ["2", "Order", "Pass"],
        ["3", "Table", "Watch"],
    ]
    conclusion = CONCLUSION if cjk else "If the text above is selectable, editable and ordered correctly, it is readable."

    story = []
    for kind, value in text:
        story.append(Paragraph(value, {"title": title, "body": body, "h2": h2}[kind]))
        if kind == "h2":
            story.append(Spacer(1, 2))

    tbl = Table(rows, colWidths=[20 * mm, 80 * mm, 30 * mm])
    tbl.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, -1), base_font),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story += [Spacer(1, 4), tbl, Spacer(1, 8)]
    story.append(Paragraph(conclusion, body))

    SimpleDocTemplate(str(path), pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm).build(story)


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="conv-test-"))
    pdf = tmp / "sample.pdf"
    docx = tmp / "sample.docx"

    print("== 准备测试夹具 ==")
    try:
        font = register_cjk()
    except ImportError:
        print("  缺少 reportlab，请先安装：pip install reportlab")
        return 2

    make_pdf(pdf, font)
    print(f"  {pdf.name}  {pdf.stat().st_size} bytes  字体={font or 'Helvetica(ASCII)'}")

    print("== 调用生产转换器 convert_worker.py ==")
    t0 = time.time()
    proc = subprocess.run(
        [sys.executable, str(WORKER), str(pdf), str(docx)],
        capture_output=True,
        timeout=300,
    )
    elapsed = time.time() - t0
    print(f"  退出码={proc.returncode}  耗时={elapsed:.1f}s")

    if proc.returncode != 0:
        print("  转换失败：")
        print("  " + (proc.stderr or b"").decode("utf-8", "replace").strip()[:800])
        return 1
    if not docx.exists() or docx.stat().st_size == 0:
        print("  未生成有效 docx")
        return 1
    print(f"  输出 {docx.name}  {docx.stat().st_size} bytes")

    print("== 校验产物可读性 ==")
    try:
        from docx import Document
    except ImportError:
        print("  （未安装 python-docx，跳过内容校验）")
        return 0

    d = Document(str(docx))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    text = "\n".join(paras)

    print(f"  段落数: {len(paras)}  表格数: {len(d.tables)}")
    for i, p in enumerate(paras[:6], 1):
        print(f"    [{i}] {p[:60]}")

    checks = [("提取到文字", len(text) > 30)]
    if font:
        checks += [
            ("包含中文关键词‘测试’", "测试" in text),
            ("标题顺序正确", bool(paras) and "报告" in paras[0]),
            ("无缺字占位符", "IIII" not in text and "\ufffd" not in text),
        ]
    else:
        checks += [("包含英文关键词 'Test'", "Test" in text)]
    checks.append(("识别到表格", len(d.tables) >= 1))

    fails = 0
    for name, ok in checks:
        print(f"  [{' OK ' if ok else 'FAIL'}] {name}")
        if not ok:
            fails += 1

    if fails:
        print(f"\n  可读性检查未全通过（{fails} 项）。产物在：{docx}")
        return 1

    print(f"\n  转换链路可用，产物：{docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
